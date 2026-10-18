"""Document Ingestion Agent (deterministic, no LLM).

Turns a deal document (PDF, DOCX, TXT, MD) into citable chunks with stable IDs like
'p2:7.2' (page 2, clause 7.2), and loads the compliance rules:
  - structured rules (YAML, JSON, CSV) are read and validated here, in code;
  - a written policy (PDF, DOCX, TXT, MD) is chunked here (ids 'pol:...') and handed to
    the Policy Parser agent, which turns it into rules.

Why no LLM here: chunk IDs are the backbone of traceability. They must be identical on
every run, so they come from code, not from a model.
"""
from __future__ import annotations

import csv
import io
import json
import re
from pathlib import Path

import pymupdf
import yaml
from pydantic import ValidationError

from src.state import AuditEntry, Chunk, DocumentInfo, PipelineState, RuleSet

# "7. Security and Collateral"  or  "SCHEDULE 1 Key Dates ..."
SECTION_RE = re.compile(r"^(?P<num>\d{1,2})\.\s+(?P<title>[A-Z\[].*)$")
SCHEDULE_RE = re.compile(r"^(?P<num>SCHEDULE\s+\d+)\s+(?P<title>.*)$", re.IGNORECASE)
# "7.2  An independent appraisal...", "S1.1  Effective Date...", "Clause 2.1 — Interest. ..."
CLAUSE_RE = re.compile(
    r"^(?:clause\s+)?(?P<num>(?:S\d+|\d{1,2})\.\d{1,2})(?:\s+[—–-]\s+|\s+)(?P<text>.+)$",
    re.IGNORECASE,
)
# Page furniture we don't want inside chunks
NOISE_RE = re.compile(r"^(Page \d+( of \d+)?|.*SYNTHETIC DOCUMENT.*)$", re.IGNORECASE)

MIN_TEXT_CHARS = 50  # below this, the PDF is probably scanned / image-only


class IngestionError(Exception):
    """Raised when the document or rule set can't be used at all."""


# ---------------------------------------------------------------- document

def _read_pages(path: Path) -> list[str]:
    suffix = path.suffix.lower()
    if suffix in {".txt", ".md"}:
        return [path.read_text(encoding="utf-8")]
    if suffix == ".docx":
        return [_read_docx(path)]
    if suffix != ".pdf":
        raise IngestionError(f"Unsupported file type: {path.suffix} (use .pdf, .docx, .txt or .md)")
    try:
        with pymupdf.open(path) as doc:
            return [page.get_text("text") for page in doc]
    except Exception as exc:  # corrupt / encrypted PDF
        raise IngestionError(f"Could not open PDF {path.name}: {exc}") from exc


def _read_docx(path: Path) -> str:
    """Word file -> plain text: paragraphs in order, then each table row as one line.

    A .docx has no fixed pages, so the whole file is treated as page 1.
    """
    try:
        import docx  # python-docx
        document = docx.Document(str(path))
    except Exception as exc:
        raise IngestionError(f"Could not open Word file {path.name}: {exc}") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("  ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _section_key(clause_num: str) -> str:
    """'7.2' -> '7', 'S1.1' -> 'SCHEDULE 1'."""
    head = clause_num.split(".")[0]
    return f"SCHEDULE {head[1:]}" if head.startswith("S") else head


MAX_NUMBER_JUMP = 3  # allow a few skipped numbers (4.1 -> 4.3) but not 4.2 -> 4.75


def _next_in_order(number: int, last: int) -> bool:
    """True if `number` can plausibly be the next section/clause after `last`."""
    return last < number <= last + MAX_NUMBER_JUMP


def _repeated_lines(pages: list[str]) -> set[str]:
    """Lines that appear on every page (running headers/footers) of a multi-page document."""
    if len(pages) < 2:
        return set()
    per_page = [{ln.strip() for ln in p.splitlines() if ln.strip()} for p in pages]
    return set.intersection(*per_page)


def _section_chunk_key(section_num: str) -> str:
    """'6' -> '6', 'SCHEDULE 1' -> 'S1', '' (before any heading) -> 'intro'."""
    if not section_num:
        return "intro"
    return section_num.replace("SCHEDULE ", "S")


def chunk_document(pages: list[str]) -> tuple[list[Chunk], str, bool]:
    """Return (chunks, title, structured)."""
    furniture = _repeated_lines(pages)
    lines: list[tuple[int, str]] = []
    for page_no, text in enumerate(pages, start=1):
        for raw in text.splitlines():
            line = raw.strip()
            if line and not NOISE_RE.match(line) and line not in furniture:
                lines.append((page_no, line))

    title = lines[0][1] if lines else "Untitled"
    chunks: list[Chunk] = []
    section_num, section_title = "", "Preamble"
    last_section = 0       # number of the last numbered section accepted
    last_clause = 0        # minor number of the last clause accepted in this section
    saw_structure = False  # any section heading or numbered clause found?
    current: dict | None = None

    def flush():
        if current and current["text"].strip():
            chunks.append(Chunk(**current))

    for page_no, line in lines:
        if m := SCHEDULE_RE.match(line):
            flush(); current = None
            section_num, section_title = m["num"].upper(), f"{m['num'].upper()} {m['title']}"
            last_clause, saw_structure = 0, True
            continue
        m = SECTION_RE.match(line)
        if m and (last_section == 0 or _next_in_order(int(m["num"]), last_section)):
            flush(); current = None
            section_num, section_title = m["num"], f"{m['num']}. {m['title']}"
            last_section, last_clause, saw_structure = int(m["num"]), 0, True
            continue
        m = CLAUSE_RE.match(line)
        # A real clause number must (1) belong to the current section and (2) come next in
        # order. So inside section 4, after clause 4.2, a wrapped line starting with
        # "2.75 per cent" (wrong section) or "4.75 per cent" (out of order) is treated as
        # a continuation, not as a new clause.
        if (m and _section_key(m["num"]) == section_num
                and _next_in_order(int(m["num"].split(".")[1]), last_clause)):
            last_clause = int(m["num"].split(".")[1])
            saw_structure = True
            flush()
            current = {
                "chunk_id": f"p{page_no}:{m['num'].upper()}",
                "page": page_no,
                "section": section_title,
                "clause": m["num"].upper(),
                "text": m["text"],
            }
        elif current is not None:
            current["text"] += " " + line  # continuation of the previous clause
        else:
            # Text that isn't inside a numbered clause: the preamble, or a section that has a
            # table or a paragraph instead of clauses. Kept as a section-level chunk
            # (e.g. 'p2:6' = page 2, section 6) so it can still be cited, never dropped.
            current = {
                "chunk_id": f"p{page_no}:{_section_chunk_key(section_num)}",
                "page": page_no,
                "section": section_title,
                "clause": None,
                "text": line,
            }

    flush()

    if saw_structure:
        return chunks, title, True

    # Fallback for documents with no numbered sections or clauses at all: paragraph
    # blocks (split further if long), still traceable by page.
    fallback: list[Chunk] = []
    for page_no, text in enumerate(pages, start=1):
        n = 0
        for block in (p for p in re.split(r"\n\s*\n", text) if p.strip()):
            for piece in _split_long(block):
                n += 1
                fallback.append(Chunk(
                    chunk_id=f"p{page_no}:b{n}", page=page_no, section=f"Page {page_no}", text=piece,
                ))
    return fallback, title, False


def _split_long(block: str, limit: int = 600) -> list[str]:
    """Split a long paragraph block into pieces of whole lines, each up to ~limit characters."""
    pieces, buf = [], ""
    for line in (ln.strip() for ln in block.splitlines() if ln.strip()):
        if buf and len(buf) + len(line) > limit:
            pieces.append(buf)
            buf = ""
        buf = f"{buf} {line}".strip()
    if buf:
        pieces.append(buf)
    return pieces


# ---------------------------------------------------------------- rules

# Rule sets written by other people use other words. These map them onto ours.
OPERATOR_ALIASES = {
    "<=": "<=", "le": "<=", "lte": "<=", "max": "<=", "at_most": "<=", "maximum": "<=",
    "less_than_or_equal": "<=", "less_than_or_equal_to": "<=", "not_more_than": "<=",
    ">=": ">=", "ge": ">=", "gte": ">=", "min": ">=", "at_least": ">=", "minimum": ">=",
    "greater_than_or_equal": ">=", "greater_than_or_equal_to": ">=", "not_less_than": ">=",
    "<": "<", "lt": "<", "less_than": "<",
    ">": ">", "gt": ">", "greater_than": ">",
    "==": "==", "=": "==", "eq": "==", "equal": "==", "equals": "==",
}
PRESENCE_OPERATORS = {"must_exist", "exists", "required", "present", "must_be_present", "must_be_stated"}


def _normalize_ruleset(data) -> dict:
    """Accept common alternative rule-set layouts and turn them into ours.

    rule_id -> id, field -> term, threshold -> value, 'less_than_or_equal' -> '<=',
    and a missing check_type is inferred: numeric comparison -> threshold,
    'must_exist' -> presence, anything else -> judgment (the operator text is added to the
    description so the LLM knows what to check).
    """
    if not isinstance(data, dict) or not isinstance(data.get("rules"), list):
        raise IngestionError("Rule set must be a YAML mapping with a 'rules:' list")
    out = dict(data)
    out["policy_name"] = str(data.get("policy_name") or data.get("name") or data.get("policy_id") or "Custom policy")
    out["version"] = str(data.get("version") or data.get("policy_id") or "1.0")

    rules = []
    for i, raw in enumerate(data["rules"], start=1):
        if not isinstance(raw, dict):
            raise IngestionError(f"Rule #{i} is not a mapping of fields")
        r = dict(raw)
        r["id"] = str(r.pop("rule_id", None) or r.get("id") or f"R-{i:02d}")
        if "term" not in r and "field" in r:
            r["term"] = r.pop("field")
        if "value" not in r and "threshold" in r:
            r["value"] = r.pop("threshold")
        raw_op = str(r.get("operator") or "").strip().lower()
        r["operator"] = OPERATOR_ALIASES.get(raw_op)
        numeric = isinstance(r.get("value"), (int, float)) and not isinstance(r.get("value"), bool)
        if "check_type" not in r:
            if r["operator"] and numeric and r.get("term"):
                r["check_type"] = "threshold"
            elif raw_op in PRESENCE_OPERATORS:
                r["check_type"] = "presence"
            else:
                r["check_type"] = "judgment"
        if not numeric:
            r.pop("value", None)
        r.setdefault("name", r["id"])
        r.setdefault("category", "general")
        r["description"] = str(r.get("description") or r["name"])
        if raw_op and not r["operator"]:
            r["description"] += f" (required check: {raw_op})"
        r["severity"] = str(r.get("severity") or "medium").lower()
        rules.append(r)
    out["rules"] = rules
    return out


STRUCTURED_RULE_FORMATS = {".yaml", ".yml", ".json", ".csv"}   # read directly by code
DOCUMENT_RULE_FORMATS = {".pdf", ".docx", ".txt", ".md"}         # read by the Policy Parser (LLM)


def _maybe_number(text: str):
    try:
        return int(text) if re.fullmatch(r"-?\d+", text) else float(text)
    except ValueError:
        return text


def _read_structured(path: Path):
    """YAML, JSON or CSV -> plain Python data (normalised afterwards)."""
    text = path.read_text(encoding="utf-8")
    suffix = path.suffix.lower()
    if suffix == ".json":
        try:
            return json.loads(text)
        except json.JSONDecodeError as exc:
            raise IngestionError(f"JSON syntax error in {path.name} at line {exc.lineno}, column {exc.colno}: {exc.msg}") from exc
    if suffix == ".csv":
        # One rule per row; the header row names the fields (id/rule_id, name, field/term, operator, ...).
        rows = []
        for row in csv.DictReader(io.StringIO(text)):
            clean = {k.strip(): _maybe_number(v.strip()) for k, v in row.items() if k and v and v.strip()}
            if clean:
                rows.append(clean)
        return {"policy_name": path.stem.replace("_", " ").title(), "rules": rows}
    return yaml.safe_load(text)


def load_rules(path: str | Path) -> RuleSet:
    """Load a STRUCTURED rule set (YAML, JSON or CSV). Documents go through load_policy()."""
    path = Path(path)
    if path.suffix.lower() not in STRUCTURED_RULE_FORMATS:
        raise IngestionError(
            f"{path.name} is not a structured rule file (.yaml, .yml, .json, .csv); "
            "policy documents are read by the Policy Parser"
        )
    try:
        data = _read_structured(path)
        ruleset = RuleSet.model_validate(_normalize_ruleset(data))
    except FileNotFoundError as exc:
        raise IngestionError(f"Rule set not found: {path}") from exc
    except yaml.YAMLError as exc:
        mark = getattr(exc, "problem_mark", None)
        where = f" at line {mark.line + 1}, column {mark.column + 1}" if mark else ""
        raise IngestionError(
            f"YAML syntax error in {path.name}{where}: {getattr(exc, 'problem', exc)}. "
            "Check the indentation: every '- ' rule entry must start at the same column."
        ) from exc
    except ValidationError as exc:
        raise IngestionError(f"Invalid rule set {path.name}: {exc}") from exc

    ids = [r.id for r in ruleset.rules]
    if len(ids) != len(set(ids)):
        raise IngestionError("Rule set contains duplicate rule IDs")
    for r in ruleset.rules:
        if r.check_type == "threshold" and (r.term is None or r.operator is None or r.value is None):
            raise IngestionError(f"Threshold rule {r.id} needs term, operator and value")
    return ruleset


def load_policy(path: str | Path) -> tuple[RuleSet | None, list[Chunk], str]:
    """Load the compliance rules in any supported format.

    Returns (ruleset, policy_chunks, source):
      - structured file (.yaml/.yml/.json/.csv) -> (RuleSet, [], "structured")
      - policy document (.pdf/.docx/.txt/.md)   -> (None, chunks, "document"); the Policy
        Parser agent turns the chunks into a RuleSet. Chunk ids get a 'pol:' prefix so a
        policy citation can never be confused with a deal citation.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in STRUCTURED_RULE_FORMATS:
        return load_rules(path), [], "structured"
    if suffix not in DOCUMENT_RULE_FORMATS:
        raise IngestionError(
            f"Unsupported rule set format: {path.suffix} "
            f"(use {', '.join(sorted(STRUCTURED_RULE_FORMATS | DOCUMENT_RULE_FORMATS))})"
        )
    if not path.exists():
        raise IngestionError(f"Rule set not found: {path}")
    pages = _read_pages(path)
    if sum(len(p.strip()) for p in pages) < MIN_TEXT_CHARS:
        raise IngestionError(f"Policy document {path.name} has no extractable text (scanned PDF?)")
    chunks, _, _ = chunk_document(pages)
    for c in chunks:
        c.chunk_id = f"pol:{c.chunk_id}"
    return None, chunks, "document"


# ---------------------------------------------------------------- agent entrypoint

def ingest(deal_path: str | Path, rules_path: str | Path) -> dict:
    """Pure function used by the orchestrator node and by tests."""
    deal_path, rules_path = Path(deal_path), Path(rules_path)
    if not deal_path.exists():
        raise IngestionError(f"Deal document not found: {deal_path}")

    pages = _read_pages(deal_path)
    if sum(len(p.strip()) for p in pages) < MIN_TEXT_CHARS:
        raise IngestionError(
            "Document has no extractable text (scanned/image PDF?). OCR is out of scope."
        )

    chunks, title, structured = chunk_document(pages)
    ruleset, policy_chunks, policy_source = load_policy(rules_path)

    info = DocumentInfo(
        path=str(deal_path), title=title, page_count=len(pages),
        chunk_count=len(chunks), structured=structured,
    )
    rules_note = (
        f"{len(ruleset.rules)} rules loaded from '{ruleset.policy_name}'" if ruleset
        else f"policy document {rules_path.name} -> {len(policy_chunks)} chunks for the Policy Parser"
    )
    audit = [AuditEntry(
        agent="ingestion", event="completed",
        detail=f"{len(pages)} pages -> {len(chunks)} chunks (structured={structured}); {rules_note}",
    )]
    return {"document": info, "chunks": chunks, "ruleset": ruleset, "policy_chunks": policy_chunks,
            "policy_source": policy_source, "audit_log": audit}


def ingestion_node(state: PipelineState) -> dict:
    """LangGraph node wrapper: converts failures into state instead of crashing the graph."""
    try:
        return ingest(state["deal_path"], state["rules_path"])
    except IngestionError as exc:
        return {
            "status": "failed",
            "errors": [f"ingestion: {exc}"],
            "audit_log": [AuditEntry(agent="ingestion", event="failed", detail=str(exc))],
        }
